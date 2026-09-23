#!/usr/bin/env python3
"""Export the Feature-MAE 32-to-64 hierarchy as a readable Excel workbook."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_ROOT = Path(
    "outputs/experiments/dinov3_multicity/feature_mae_n30x12800_qc/mae/"
    "hierarchy_32_64"
)
DEFAULT_OUTPUT = Path(
    "paper/tables/supplementary/Table_MAE_Category_Evolution_32_to_64.xlsx"
)
DEFAULT_DATA_DIR = Path("paper/data/feature_mae_hierarchy")
SITE = "https://cehougis.github.io/urban-visual-gene/mae.html"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def build_frames(root: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    taxonomy = read_json(root / "taxonomy.json")
    report = read_json(root / "hierarchy_report.json")
    labels = read_json(root / "semantic_labels_qwen" / "category_labels.json")
    names = labels["categories"]
    feature_rows = pd.read_csv(root / "feature_hierarchy.csv")

    transitions = []
    coarse_rows = []
    for coarse in taxonomy["children"]:
        coarse_id = coarse["id"]
        coarse_name = names[coarse_id]
        children = coarse["children"]
        child_names = "；".join(
            f"{child['id']} {names[child['id']]['name_zh']}（{child['n']}维）"
            for child in children
        )
        coarse_rows.append({
            "粗类编号": coarse_id,
            "粗类名称": coarse_name["name_zh"],
            "英文名称": coarse_name["name_en"],
            "维度数": coarse["n"],
            "细类数": len(children),
            "统计中心维度": coarse["medoid_feature"],
            "统计稳定性": coarse["stability"],
            "Qwen命名置信度": coarse_name["confidence"],
            "视觉总结": coarse_name["summary_zh"],
            "包含细类": child_names,
            "成员维度": ", ".join(map(str, sorted(
                feature for child in children for feature in child["features"]
            ))),
            "中心维度网页": f"{SITE}#dim-{coarse['medoid_feature']:03d}",
        })
        for fine in children:
            fine_id = fine["id"]
            fine_name = names[fine_id]
            ratio = fine["n"] / coarse["n"]
            transitions.append({
                "粗类编号": coarse_id,
                "粗类名称": coarse_name["name_zh"],
                "粗类维度数": coarse["n"],
                "粗类统计稳定性": coarse["stability"],
                "包含细类数": len(children),
                "细类编号": fine_id,
                "细类名称": fine_name["name_zh"],
                "细类英文名": fine_name["name_en"],
                "细类维度数": fine["n"],
                "占父粗类比例": ratio,
                "细类统计稳定性": fine["stability"],
                "Qwen命名置信度": fine_name["confidence"],
                "语义类型": fine_name["semantic_type"],
                "统计中心维度": fine["medoid_feature"],
                "成员维度": ", ".join(map(str, sorted(fine["features"]))),
                "细类视觉总结": fine_name["summary_zh"],
                "变化说明": (
                    f"{coarse_id} {coarse_name['name_zh']} → "
                    f"{fine_id} {fine_name['name_zh']}（{fine['n']}/{coarse['n']}维，"
                    f"{ratio:.1%}）"
                ),
                "中心维度网页": f"{SITE}#dim-{fine['medoid_feature']:03d}",
            })

    transitions_frame = pd.DataFrame(transitions).sort_values(
        ["粗类编号", "细类编号"]
    )
    coarse_frame = pd.DataFrame(coarse_rows).sort_values("粗类编号")
    feature_frame = feature_rows.copy()
    feature_frame["粗类名称"] = feature_frame["coarse_cluster"].map(
        lambda x: names[x]["name_zh"]
    )
    feature_frame["细类名称"] = feature_frame["fine_cluster"].map(
        lambda x: names[x]["name_zh"]
    )
    feature_frame["维度网页"] = feature_frame["feature_id"].map(
        lambda x: f"{SITE}#dim-{int(x):03d}"
    )
    feature_frame = feature_frame.rename(columns={
        "feature_id": "维度编号",
        "coarse_cluster": "粗类编号",
        "fine_cluster": "细类编号",
        "coarse_stability": "维度粗类稳定性",
        "fine_stability": "维度细类稳定性",
        "top1_patch_support": "Top1 Patch支持数",
        "image_support": "图像支持数",
    })[[
        "维度编号", "粗类编号", "粗类名称", "细类编号", "细类名称",
        "维度粗类稳定性", "维度细类稳定性", "Top1 Patch支持数",
        "图像支持数", "维度网页",
    ]].sort_values("维度编号")

    metrics = [
        {"项目": "层级关系", "数值": "32个粗类 → 64个细类，同一棵Ward树", "说明": "每个细类严格属于一个粗类"},
        {"项目": "潜变量维度", "数值": 512, "说明": "全部维度均有唯一粗类和细类归属"},
        {"项目": "聚类方法", "数值": report["method"], "说明": "语义名称不参与聚类"},
        {"项目": "相似性视图", "数值": "；".join(report["views"]), "说明": "五视图秩融合"},
        {"项目": "粗类平均半样本Jaccard", "数值": report["split_sample"]["coarse_mean_jaccard"], "说明": "越高表示半样本复现越稳定"},
        {"项目": "细类平均半样本Jaccard", "数值": report["split_sample"]["fine_mean_jaccard"], "说明": "越高表示半样本复现越稳定"},
        {"项目": "粗类NMI", "数值": report["split_sample"]["coarse_nmi"], "说明": "半样本聚类与全样本聚类的一致性"},
        {"项目": "细类NMI", "数值": report["split_sample"]["fine_nmi"], "说明": "半样本聚类与全样本聚类的一致性"},
        {"项目": "命名模型", "数值": labels["model"], "说明": "训练后后验解释"},
        {"项目": "每类命名图像数", "数值": labels["examples_per_category"], "说明": "统计中心维度的Top-10代表图"},
        {"项目": "网页", "数值": SITE, "说明": "点击维度卡片查看热图"},
    ]
    return transitions_frame, coarse_frame, feature_frame, pd.DataFrame(metrics)


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    group_fills = [
        PatternFill("solid", fgColor="EAF2F8"),
        PatternFill("solid", fgColor="FDF2E9"),
    ]
    widths = {
        "类别变化": [12, 18, 13, 16, 12, 12, 22, 20, 13, 15, 16, 17, 12, 14, 34, 42, 52, 40],
        "粗类汇总": [12, 20, 24, 10, 10, 14, 14, 17, 45, 55, 50, 40],
        "维度归属": [10, 12, 20, 12, 22, 16, 16, 18, 14, 40],
        "方法说明": [28, 70, 48],
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for index, width in enumerate(widths[sheet.title], 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if sheet.title in {"类别变化", "粗类汇总", "维度归属"}:
            coarse_col = 1 if sheet.title != "维度归属" else 2
            coarse_values = []
            for row_index in range(2, sheet.max_row + 1):
                coarse = sheet.cell(row_index, coarse_col).value
                if coarse not in coarse_values:
                    coarse_values.append(coarse)
                fill = group_fills[coarse_values.index(coarse) % 2]
                for cell in sheet[row_index]:
                    cell.fill = fill
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("https://"):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    args = parser.parse_args()
    root = args.root.resolve()
    output = args.output.resolve()
    data_dir = args.data_dir.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    data_dir.mkdir(parents=True, exist_ok=True)
    transitions, coarse, features, methods = build_frames(root)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        transitions.to_excel(writer, sheet_name="类别变化", index=False)
        coarse.to_excel(writer, sheet_name="粗类汇总", index=False)
        features.to_excel(writer, sheet_name="维度归属", index=False)
        methods.to_excel(writer, sheet_name="方法说明", index=False)
    style_workbook(output)
    transitions.to_csv(
        data_dir / "category_transitions_32_to_64.csv",
        index=False, encoding="utf-8-sig",
    )
    coarse.to_csv(
        data_dir / "coarse_categories_32.csv",
        index=False, encoding="utf-8-sig",
    )
    features.to_csv(
        data_dir / "dimension_membership_512.csv",
        index=False, encoding="utf-8-sig",
    )
    methods.to_csv(
        data_dir / "hierarchy_method_metrics.csv",
        index=False, encoding="utf-8-sig",
    )
    print(f"xlsx={output}")
    print(f"data_dir={data_dir}")
    print(f"transitions={len(transitions)} coarse={len(coarse)} features={len(features)}")


if __name__ == "__main__":
    main()
